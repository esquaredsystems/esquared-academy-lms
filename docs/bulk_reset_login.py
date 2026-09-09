"""
Bulk-run `manage.py reset_login <ID> --set` for every student row in the Excel file,
feeding the password (from column C, parentheses stripped) into the command's
"New password:" / confirm prompts.

ID   -> column E   (e.g. 26070108)
PW   -> column C   (stored as "(Aariz)"; parentheses are stripped -> "Aariz")

Setup (run once in your project folder):
    venv/Scripts/pip install openpyxl

Usage:
    1. TEST FIRST: with END_ROW = 3 and DRY_RUN = False, run once, then log in as
       26070108 to confirm the password is what you expect.
    2. Then set END_ROW = None and run again to process all students.

    py bulk_reset_login.py
"""

import subprocess
import sys
from pathlib import Path
from openpyxl import load_workbook

# ---------------- CONFIG ----------------
EXCEL_PATH  = r"C:\Users\uzair\OneDrive\Documents\GitHub\esquared-academy-lms\docs\Esquared_Academy_Logins.xlsx"
SHEET_NAME  = "Students"      # which sheet to process (Students / Teachers / Examiners / Other Roles)
ID_COL      = "E"             # column with the IDs
PW_COL      = "C"             # column with the passwords
START_ROW   = 3               # first data row (row 1 = note, row 2 = header)
END_ROW     = None               # last row; keep 3 to test ONE row, set None for all rows

STRIP_PARENS = True           # True -> "(Aariz)" becomes "Aariz"

PROJECT_DIR = r"C:\Users\uzair\OneDrive\Documents\GitHub\esquared-academy-lms"
PYTHON      = str(Path(PROJECT_DIR) / "venv" / "Scripts" / "python.exe")

DRY_RUN     = False            # True = only print commands; False = actually run them
# ----------------------------------------


def build_command(user_id: str) -> list:
    # --set is a flag; the password is typed into the prompt, not passed here.
    return [PYTHON, "manage.py", "reset_login", user_id, "--set"]


def clean_id(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def clean_password(value) -> str:
    if value is None:
        return ""
    s = str(value).strip()
    if STRIP_PARENS and s.startswith("(") and s.endswith(")"):
        s = s[1:-1].strip()
    return s


def main() -> None:
    wb = load_workbook(EXCEL_PATH, data_only=True)
    ws = wb[SHEET_NAME] if SHEET_NAME else wb.active

    last_row = END_ROW if END_ROW else ws.max_row
    ok, failed, skipped = 0, 0, 0

    for row in range(START_ROW, last_row + 1):
        user_id  = clean_id(ws[f"{ID_COL}{row}"].value)
        password = clean_password(ws[f"{PW_COL}{row}"].value)

        if not user_id or not password:
            skipped += 1
            continue

        cmd = build_command(user_id)

        if DRY_RUN:
            print(f"row {row}:  ID {user_id}   password -> {password!r}")
            continue

        try:
            # password fed twice: the command prompts to type AND confirm it.
            result = subprocess.run(
                cmd,
                cwd=PROJECT_DIR,
                capture_output=True,
                text=True,
                input=(password + "\n") * 2,
                timeout=30,          # if the prompt ignores piped input, fail here instead of hanging
            )
            if result.returncode == 0:
                ok += 1
                print(f"[OK]   row {row}  ID {user_id}")
            else:
                failed += 1
                print(f"[FAIL] row {row}  ID {user_id}  ->  {result.stderr.strip() or result.stdout.strip()}")
        except subprocess.TimeoutExpired:
            failed += 1
            print(f"[HANG] row {row}  ID {user_id}  ->  prompt did not read the piped password (see note)")
        except Exception as e:  # noqa: BLE001
            failed += 1
            print(f"[ERR]  row {row}  ID {user_id}  ->  {e}")

    print("\n----- summary -----")
    if DRY_RUN:
        print("DRY RUN only -- nothing was executed. Set DRY_RUN = False to run.")
    else:
        print(f"success: {ok}   failed: {failed}   skipped(blank): {skipped}")


if __name__ == "__main__":
    if not Path(EXCEL_PATH).exists():
        sys.exit(f"Excel file not found: {EXCEL_PATH}")
    main()
