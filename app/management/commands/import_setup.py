"""
Load the filled-in Esquared_LMS_Setup workbook into the database.

    python manage.py import_setup docs/Esquared_LMS_Setup.xlsx --dry-run
    python manage.py import_setup docs/Esquared_LMS_Setup.xlsx

Safe to run more than once: everything is matched on a natural key and
updated rather than duplicated, so a corrected workbook can simply be
re-imported.

Staff accounts are created WITHOUT a password (Django's "unusable
password"). Nothing here can log in until someone sets a password:

    python manage.py reset_login <username> --set

That is deliberate — passwords are never written by a script, put on a
command line, or left in shell history.
"""

import json
import re
from datetime import date, datetime, time, timedelta
from pathlib import Path

from django.contrib.auth.models import Group
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils.text import slugify

from app import access, models

_STUDENT_ID_TYPE_MAP = {"cnic": "cnic", "b-form": "b_form", "bform": "b_form", "passport": "passport"}


def set_student_attribute(student, short_name, raw_value):
    """
    Set (or clear) a StudentAttribute by its type's short_name.

    Replaces the direct field assignment this helper's callers used before
    national_id_type/guardian_contact_2 became attributes. The attribute
    type must already exist (see the seed_attribute_types command) — a
    workbook value for a type that isn't set up is silently skipped rather
    than raising, since import_setup runs unattended.
    """
    attribute_type = models.StudentAttributeType.objects.filter(short_name=short_name).first()
    if attribute_type is None:
        return
    existing = models.StudentAttribute.objects.filter(
        student=student, attribute_type=attribute_type
    ).first()
    if not raw_value:
        if existing:
            existing.void(reason="cleared by import_setup")
        return
    if existing:
        existing.value_reference = raw_value
        existing.save()
    else:
        models.StudentAttribute.objects.create(
            student=student, attribute_type=attribute_type, value_reference=raw_value,
        )

#: The school year runs 1 July 2026 – 30 June 2027 and is called "2627".
DEFAULT_YEAR = 2627
YEAR_START = date(2026, 7, 1)

#: Admission numbers are YYMMDDHH — the date and hour a student was
#: admitted. These students pre-date the system, so the initial load walks
#: forward one hour per student in alphabetical order from the start of the
#: school year. They are placeholders: real numbers get set when the actual
#: joining dates are known.
ADMISSION_EPOCH = datetime(2026, 7, 1, 8, 0)

ACADEMY_CLASS_LEVELS = {"E1": 1, "E2": 2, "S1": 3, "S2": 4, "S3": 5}

#: What people type in the workbook -> the subject's short_name.
SUBJECT_ALIASES = {
    "english": "ENG", "eng": "ENG", "english language": "ENG",
    "math": "MATH", "maths": "MATH", "mathematics": "MATH",
    "physics": "PHY", "chemistry": "CHEM",
    "islamiyat": "ISL",
    "history": "PST1", "history of pak": "PST1",
    "env. of pak": "PST2", "env of pak": "PST2", "environment of pakistan": "PST2",
    "computer science": "CS", "cs": "CS",
    "business": "BUS", "business studies": "BUS",
    "accounts": "ACC", "accounting": "ACC",
    "economics": "ECON",
    "urdu": "URDU", "urdu language": "URDU",
    "activity": "ACT",
}
#: Subjects that are not on the reference sheet but are taught anyway.
EXTRA_SUBJECTS = {"ACT": "Activity"}

DAYS = {"monday": 0, "tuesday": 1, "wednesday": 2, "thursday": 3,
        "friday": 4, "saturday": 5, "sunday": 6}


# ---------------------------------------------------------------------
# Reading the workbook
#
# The .xlsx path needs openpyxl, which is not part of this project's
# requirements. So the same data can also be supplied as a .json file
# (sheet name -> list of rows), which needs nothing extra. Both are read
# through the same tiny adapter so the import logic below never has to
# care which one it was given.
# ---------------------------------------------------------------------
class _Cell:
    __slots__ = ("value",)

    def __init__(self, value):
        self.value = value


class _Sheet:
    def __init__(self, rows):
        self._rows = rows

    def iter_rows(self, min_row=1, max_row=None):
        end = len(self._rows) if max_row is None else min(max_row, len(self._rows))
        for i in range(min_row - 1, end):
            yield [_Cell(v) for v in self._rows[i]]

    @property
    def max_row(self):
        return len(self._rows)


class _Book:
    def __init__(self, data):
        self._data = data
        self.sheetnames = list(data)

    def __getitem__(self, name):
        return _Sheet(self._data[name])


def load_book(path):
    """A workbook from .json (no dependencies) or .xlsx (needs openpyxl)."""
    if path.suffix.lower() == ".json":
        with open(path, encoding="utf-8") as fh:
            return _Book(json.load(fh))
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise CommandError(
            "Reading .xlsx needs openpyxl, which is not installed.\n"
            "Either install it:      pip install openpyxl\n"
            "or import the .json:    manage.py import_setup docs/setup_data.json"
        )
    wb = load_workbook(path, data_only=True)
    return _Book({
        name: [[c.value for c in row] for row in wb[name].iter_rows()]
        for name in wb.sheetnames
    })


def to_time(value):
    """A time from a real time object or an 'HH:MM' / 'HH:MM:SS' string."""
    if value in (None, ""):
        return None
    if isinstance(value, time):
        return value
    if hasattr(value, "hour") and hasattr(value, "minute"):
        return time(value.hour, value.minute)
    text = str(value).strip()
    for fmt in ("%H:%M", "%H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).time()
        except ValueError:
            continue
    return None


def clean(value):
    """Trim, collapse non-breaking spaces, and turn blanks into ''. """
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ").strip()
    while "  " in text:
        text = text.replace("  ", " ")
    return text


def is_example(row_values):
    """The template ships example rows; they must never be imported."""
    return any("example row" in clean(v).lower() for v in row_values)


class Command(BaseCommand):
    help = "Import the Esquared_LMS_Setup workbook (students, staff, timetable)."

    def add_arguments(self, parser):
        parser.add_argument("workbook", nargs="?",
                            default="docs/setup_data.json")
        parser.add_argument("--year", type=int, default=DEFAULT_YEAR,
                            help="Academic year to file everything under.")
        parser.add_argument("--dry-run", action="store_true",
                            help="Report what would happen, write nothing.")

    # -- helpers -------------------------------------------------------
    def subject_code(self, typed):
        key = clean(typed).lower().rstrip(".")
        code = SUBJECT_ALIASES.get(key)
        if code is None:
            self.warn(f"subject {typed!r} is not recognised — row skipped")
        return code

    def warn(self, msg):
        self.warnings.append(msg)

    def handle(self, *args, **opts):
        path = Path(opts["workbook"])
        if not path.exists():
            raise CommandError(f"No workbook at {path}")

        self.year = opts["year"]
        self.dry = opts["dry_run"]
        self.warnings = []
        self.counts = {}

        wb = load_book(path)
        for needed in ("Students", "Teachers", "Examiners", "Timetable",
                       "Grades", "Subjects"):
            if needed not in wb.sheetnames:
                raise CommandError(f"The workbook has no {needed!r} sheet.")

        self.stdout.write("")
        self.stdout.write(self.style.MIGRATE_HEADING(
            f"Importing {path.name} into academic year {self.year}"
            + ("  (DRY RUN — nothing will be written)" if self.dry else "")
        ))

        with transaction.atomic():
            self.academy_classes = self.do_academy_classes(wb)
            self.subjects = self.do_subjects(wb)
            self.teachers = self.do_teachers(wb)
            self.do_examiners(wb)
            self.do_students(wb)
            self.do_timetable(wb)
            if self.dry:
                transaction.set_rollback(True)

        self.report()

    # -- reference -----------------------------------------------------
    def do_academy_classes(self, wb):
        found = {}
        for row in wb["Grades"].iter_rows(min_row=5):
            code, full = clean(row[0].value), clean(row[1].value)
            if not code or code.lower() == "code":
                continue
            academy_class = models.AcademyClass.all_objects.filter(short_name=code).first()
            if academy_class is None:
                academy_class = models.AcademyClass(short_name=code)
            academy_class.full_name = full or code
            academy_class.level = ACADEMY_CLASS_LEVELS.get(code, academy_class.level or 0)
            academy_class.sort_order = ACADEMY_CLASS_LEVELS.get(code, 0)
            academy_class.voided = False
            academy_class.active_flag = True
            academy_class.save()
            found[code] = academy_class
        self.counts["classes"] = len(found)
        return found

    def do_subjects(self, wb):
        found = {}
        rows = []
        for row in wb["Subjects"].iter_rows(min_row=5):
            code, full, cam = (clean(row[0].value), clean(row[1].value),
                               clean(row[2].value))
            if not code or code.lower() == "code":
                continue
            rows.append((code, full, cam))
        for code, full in EXTRA_SUBJECTS.items():
            if code not in [r[0] for r in rows]:
                rows.append((code, full, ""))
                self.warn(f"subject {full!r} was not on the Subjects sheet — added it")
        for order, (code, full, cam) in enumerate(rows):
            subject = models.Subject.all_objects.filter(short_name=code).first()
            if subject is None:
                subject = models.Subject(short_name=code)
            subject.full_name = full or code
            subject.id_number = cam or None
            subject.sort_order = order
            subject.voided = False
            subject.active_flag = True
            subject.save()
            found[code] = subject
        self.counts["subjects"] = len(found)
        return found

    def syllabus_for(self, academy_class_code, subject_code):
        academy_class = self.academy_classes.get(academy_class_code)
        subject = self.subjects.get(subject_code)
        if academy_class is None:
            self.warn(f"unknown class {academy_class_code!r} — row skipped")
            return None
        if subject is None:
            return None
        syllabus = models.Syllabus.all_objects.filter(
            academy_class=academy_class, subject=subject, academic_year=self.year
        ).first()
        if syllabus is None:
            syllabus = models.Syllabus(
                academy_class=academy_class, subject=subject, academic_year=self.year
            )
        syllabus.full_name = f"{subject.full_name} — {academy_class.short_name} ({self.year})"
        syllabus.voided = False
        syllabus.active_flag = True
        syllabus.save()
        return syllabus

    # -- people --------------------------------------------------------
    def username_for(self, first, last, email):
        """
        A predictable username: the email's local part, else the name.

        Dots are kept (uxair.ahm, not uxairahm) so the login is the one a
        person would guess. A clash with a different person gets a number.
        """
        if email:
            base = email.split("@")[0].strip().lower()
        else:
            base = f"{first} {last}".strip().lower().replace(" ", ".")
        base = re.sub(r"[^a-z0-9._]", "", base).strip("._") or "staff"
        name, n = base, 1
        while models.AppUser.all_objects.filter(username=name).exists():
            n += 1
            name = f"{base}{n}"
        return name

    def staff_user(self, first, last, email, group_name):
        """One AppUser per person, in the given role group. No password."""
        user = None
        if email:
            user = models.AppUser.all_objects.filter(email=email).first()
        if user is None:
            user = models.AppUser.all_objects.filter(
                first_name=first, last_name=last
            ).first()
        if user is None:
            user = models.AppUser(username=self.username_for(first, last, email))
            user.set_unusable_password()
        user.first_name = first
        user.last_name = last
        if email:
            user.email = email
        user.is_staff = True
        user.is_active = True
        user.voided = False
        user.active_flag = True
        user.save()
        group = Group.objects.filter(name=group_name).first()
        if group is None:
            self.warn(f"role {group_name!r} does not exist — run seed_roles first")
        else:
            user.groups.add(group)
        return user

    def do_teachers(self, wb):
        teachers, seen_rows = {}, 0
        for row in wb["Teachers"].iter_rows(min_row=6):
            v = [c.value for c in row]
            if all(clean(x) == "" for x in v):
                continue
            if is_example(v):
                continue
            staff_no = clean(v[0])
            first, last = clean(v[1]), clean(v[2])
            email, mobile = clean(v[3]), clean(v[4])
            academy_class_code, subject_typed = clean(v[5]), clean(v[6])
            also_examines = clean(v[7]).lower().startswith("y")
            if not (staff_no and first):
                continue
            seen_rows += 1

            if staff_no not in teachers:
                user = self.staff_user(first, last, email, access.TEACHER)
                teacher = models.Teacher.all_objects.filter(staff_no=staff_no).first()
                if teacher is None:
                    teacher = models.Teacher(staff_no=staff_no)
                teacher.user = user
                teacher.voided = False
                teacher.active_flag = True
                teacher.save()
                teachers[staff_no] = teacher
                if also_examines:
                    group = Group.objects.filter(name=access.EXAMINER).first()
                    if group:
                        user.groups.add(group)

            code = self.subject_code(subject_typed)
            if not code:
                continue
            syllabus = self.syllabus_for(academy_class_code, code)
            if syllabus is None:
                continue
            link = models.TeachingAssignment.all_objects.filter(
                teacher=teachers[staff_no], syllabus=syllabus, role="primary"
            ).first()
            if link is None:
                link = models.TeachingAssignment(
                    teacher=teachers[staff_no], syllabus=syllabus, role="primary"
                )
            link.voided = False
            link.active_flag = True
            link.save()

        self.counts["teachers"] = len(teachers)
        self.counts["teaching rows"] = seen_rows
        return teachers

    def do_examiners(self, wb):
        """
        The Examiners sheet names the same people as the Teachers sheet.

        Its staff-number column was left as X-001 on every row, so people
        are matched by name instead, and the role is simply added to the
        account they already have.
        """
        by_name, rows = {}, 0
        for row in wb["Examiners"].iter_rows(min_row=6):
            v = [c.value for c in row]
            if all(clean(x) == "" for x in v):
                continue
            first, last = clean(v[1]), clean(v[2])
            if not first:
                continue
            rows += 1
            key = (first.lower(), last.lower())
            if key in by_name:
                continue
            user = self.staff_user(first, last, clean(v[3]), access.EXAMINER)
            by_name[key] = user
        self.counts["examiners"] = len(by_name)
        self.counts["examiner rows"] = rows

    def do_students(self, wb):
        gathered = []
        for row in wb["Students"].iter_rows(min_row=6):
            v = [c.value for c in row]
            if all(clean(x) == "" for x in v):
                continue
            if is_example(v):
                continue
            first = clean(v[1])
            if not first:
                continue
            gathered.append({
                "admission_no": clean(v[0]),
                "first_name": first,
                "last_name": clean(v[2]),
                "academy_class": clean(v[3]),
                "dob": v[4],
                "email": clean(v[5]),
                "mobile": clean(v[6]),
                "national_id": clean(v[7]),
                "id_type": clean(v[8]),
                "guardian_name": clean(v[9]),
                "guardian_contact": clean(v[10]),
                "guardian_contact_2": clean(v[11]),
                "address": clean(v[12]),
            })

        # Alphabetical across the whole school, then one hour per student.
        gathered.sort(key=lambda s: s["first_name"].lower())

        made = 0
        for i, s in enumerate(gathered):
            student = models.Student.all_objects.filter(
                first_name=s["first_name"], last_name=s["last_name"] or ""
            ).first()
            if student is None:
                student = models.Student()
            if not student.admission_no:
                stamp = ADMISSION_EPOCH + timedelta(hours=i)
                student.admission_no = s["admission_no"] or stamp.strftime("%y%m%d%H")
            student.first_name = s["first_name"]
            student.last_name = s["last_name"]
            if s["dob"]:
                try:
                    student.date_of_birth = (
                        s["dob"].date() if hasattr(s["dob"], "date")
                        else datetime.strptime(str(s["dob"])[:10], "%Y-%m-%d").date()
                    )
                except (ValueError, TypeError):
                    self.warn(f"{s['first_name']}: date of birth {s['dob']!r} not understood")
            student.email = s["email"] or None
            student.mobile = s["mobile"]
            student.address = s["address"]
            student.national_id = s["national_id"] or None
            student.guardian_name = s["guardian_name"] or None
            student.guardian_contact = s["guardian_contact"] or None
            student.voided = False
            student.active_flag = True
            student.save()
            if s["id_type"]:
                set_student_attribute(
                    student, "national_id_type", _STUDENT_ID_TYPE_MAP.get(s["id_type"].lower(), "")
                )
            set_student_attribute(student, "guardian_contact_2", s["guardian_contact_2"])
            made += 1

            academy_class = self.academy_classes.get(s["academy_class"])
            if academy_class is None:
                self.warn(f"{s['first_name']}: unknown class {s['academy_class']!r} — not enrolled")
                continue
            enrolment = models.Enrolment.all_objects.filter(
                student=student, academic_year=self.year
            ).first()
            if enrolment is None:
                enrolment = models.Enrolment(student=student, academic_year=self.year)
            enrolment.academy_class = academy_class
            enrolment.started_on = YEAR_START
            enrolment.status = "active"
            enrolment.voided = False
            enrolment.active_flag = True
            enrolment.save()

        self.counts["students"] = made

    def do_timetable(self, wb):
        made = 0
        for row in wb["Timetable"].iter_rows(min_row=6):
            v = [c.value for c in row]
            if all(clean(x) == "" for x in v):
                continue
            if is_example(v):
                continue
            academy_class_code, day_typed, period = clean(v[0]), clean(v[1]), clean(v[2])
            start, end = v[3], v[4]
            subject_typed, staff_no, room = clean(v[5]), clean(v[6]), clean(v[7])
            if not (academy_class_code and day_typed and subject_typed):
                continue
            day = DAYS.get(day_typed.lower())
            if day is None:
                self.warn(f"day {day_typed!r} not understood — row skipped")
                continue
            code = self.subject_code(subject_typed)
            if not code:
                continue
            syllabus = self.syllabus_for(academy_class_code, code)
            if syllabus is None:
                continue
            teacher = self.teachers.get(staff_no)
            if staff_no and teacher is None:
                self.warn(f"timetable names staff {staff_no!r}, who is not on the "
                          "Teachers sheet — slot left without a teacher")

            slot = models.TimetableSlot.all_objects.filter(
                syllabus=syllabus, day_of_week=day, period=period or "1"
            ).first()
            if slot is None:
                slot = models.TimetableSlot(
                    syllabus=syllabus, day_of_week=day, period=period or "1"
                )
            slot.teacher = teacher
            slot.start_time = to_time(start)
            slot.end_time = to_time(end)
            slot.voided = False
            slot.active_flag = True
            slot.save()
            made += 1
        self.counts["timetable slots"] = made

    # -- output --------------------------------------------------------
    def report(self):
        self.stdout.write("")
        for label, n in self.counts.items():
            self.stdout.write(f"  {label:<18} {n}")
        if self.warnings:
            self.stdout.write("")
            self.stdout.write(self.style.WARNING("Warnings"))
            for w in dict.fromkeys(self.warnings):
                self.stdout.write(f"  - {w}")
        self.stdout.write("")
        if self.dry:
            self.stdout.write(self.style.WARNING(
                "DRY RUN — nothing was written. Re-run without --dry-run to apply."))
        else:
            self.stdout.write(self.style.SUCCESS("Imported."))
            self.stdout.write(
                "Staff accounts have NO password yet. For each one:\n"
                "    manage.py reset_login <username> --set"
            )
